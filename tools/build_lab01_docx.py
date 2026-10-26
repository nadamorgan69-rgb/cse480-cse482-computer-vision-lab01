from pathlib import Path
import json

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPOSITORY_URL = "https://github.com/nadamorgan69-rgb/cse480-cse482-computer-vision-lab01"
RESULTS = ROOT / "results"
COVER_IMAGE = ROOT / "images" / "computer_vision_dream_cover.png"
OUTPUT = ROOT / "report" / "Lab01_Report.docx"
DATA = json.loads((RESULTS / "verification.json").read_text(encoding="utf-8"))
METRICS = DATA["metrics"]

NAVY = "1F4E78"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9D9D9"
PALE_GRAY = "F5F7F9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def style_table(table, widths, numeric_columns=()):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_layout(table)
    set_table_borders(table)
    mark_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
            )
            for run in paragraph.runs:
                run.font.name = "Arial"
                run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else BLACK
                run.bold = row_index == 0


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(doc, relative_path, caption, width=Inches(7.0), alt_text=""):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(2)
    shape = paragraph.add_run().add_picture(str(RESULTS / relative_path), width=width)
    set_alt_text(shape, caption, alt_text or caption)
    caption_paragraph = doc.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = True


def add_labeled_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.keep_together = True
    lead = paragraph.add_run(label + " ")
    lead.bold = True
    paragraph.add_run(text)
    return paragraph


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_native_equation(doc, equation_text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    math_para = OxmlElement("m:oMathPara")
    math_obj = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = equation_text
    math_run.append(math_text)
    math_obj.append(math_run)
    math_para.append(math_obj)
    paragraph._p.append(math_para)


def add_page_break(doc):
    doc.add_page_break()


def add_task_heading(doc, text):
    heading = doc.add_heading(text, level=1)
    heading.paragraph_format.page_break_before = True
    return heading


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("CSE480 CSE482   Computer Vision Lab   Assignment 01   Page ")
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    for r in paragraph.runs:
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(89, 89, 89)


def remove_paragraph_borders(property_container):
    borders = property_container.find(qn("w:pBdr"))
    if borders is not None:
        property_container.remove(borders)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.30)


def restart_page_numbering(section, start=1):
    sect_pr = section._sectPr
    page_number_type = sect_pr.find(qn("w:pgNumType"))
    if page_number_type is None:
        page_number_type = OxmlElement("w:pgNumType")
        sect_pr.append(page_number_type)
    page_number_type.set(qn("w:start"), str(start))


doc = Document()
section = doc.sections[0]
configure_section(section)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(12)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(18)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(6)
title_style_ppr = title_style._element.find(qn("w:pPr"))
if title_style_ppr is None:
    title_style_ppr = OxmlElement("w:pPr")
    title_style._element.append(title_style_ppr)
remove_paragraph_borders(title_style_ppr)

for name, size, before, after in (("Heading 1", 18, 7, 6), ("Heading 2", 16, 6, 4)):
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

caption_style = styles["Caption"]
caption_style.font.name = "Arial"
caption_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
caption_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
caption_style.font.size = Pt(12)
caption_style.font.italic = True
caption_style.font.color.rgb = RGBColor(80, 80, 80)
caption_style.paragraph_format.space_after = Pt(5)

# Cover page
cover_spacer = doc.add_paragraph()
cover_spacer.paragraph_format.space_after = Pt(8)
title = doc.add_paragraph("Computer Vision Lab Assignment 01", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
remove_paragraph_borders(title._p.get_or_add_pPr())
subtitle = doc.add_paragraph("CSE480 CSE482 Computer Vision Lab")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True
subtitle.runs[0].font.size = Pt(16)
program = doc.add_paragraph("Ain Shams University  Mechatronics Engineering and Automation Program")
program.alignment = WD_ALIGN_PARAGRAPH.CENTER
program.paragraph_format.space_after = Pt(12)
program.runs[0].font.size = Pt(12)

cover_visual = doc.add_paragraph()
cover_visual.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover_visual.paragraph_format.space_after = Pt(12)
cover_shape = cover_visual.add_run().add_picture(str(COVER_IMAGE), width=Inches(6.85))
set_alt_text(
    cover_shape,
    "Dreamy computer vision cover artwork",
    "A luminous camera lens and abstract eye surrounded by neural networks, pixel grids, feature points, contours, and depth-map layers in blue and violet light.",
)

student_name = doc.add_paragraph()
student_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
student_name.paragraph_format.space_after = Pt(4)
name_label = student_name.add_run("Student name  ")
name_label.bold = True
student_name.add_run("Nada Said Yassin Morgan")
for run in student_name.runs:
    run.font.size = Pt(12)

student_id = doc.add_paragraph()
student_id.alignment = WD_ALIGN_PARAGRAPH.CENTER
id_label = student_id.add_run("Student ID  ")
id_label.bold = True
student_id.add_run("23P0284")
for run in student_id.runs:
    run.font.size = Pt(12)

content_section = doc.add_section(WD_SECTION.NEW_PAGE)
configure_section(content_section)
content_section.footer.is_linked_to_previous = False
add_page_number(content_section.footer.paragraphs[0])
restart_page_numbering(content_section, 1)

doc.add_heading("How I approached the lab", level=2)
doc.add_paragraph(
    "I treated the four tasks as small experiments in how pixel values control an image. I kept the supplied "
    "images unchanged, processed their NumPy arrays with OpenCV, and converted BGR to RGB only for display. "
    "For each result, I relate the visible change to the mask, weights, or channel assignments that produced it."
)

doc.add_heading("Task 1 Selective Colour Isolation", level=1)
add_labeled_paragraph(
    doc,
    "Reasoning and method",
    "I converted the apple image to grayscale so the unselected region would keep its brightness and texture "
    "without colour. I split the original into blue, green, and red planes, remembering that OpenCV returns BGR. "
    "A large red value alone can also describe yellow pixels, so the mask requires R >= 70, R - G >= 30, "
    "R - B >= 25, and G - B <= 35. Signed subtraction avoids uint8 wraparound. Tighter red gaps would select "
    "fewer, purer pixels; looser gaps would keep more skin but also more orange and yellow areas."
)
add_figure(
    doc,
    "task1/stages.png",
    "Figure 1  Original image grayscale red channel binary mask extracted colours and final composition",
    width=Inches(6.35),
    alt_text="Original apples, grayscale image, red channel, binary red mask, extracted original colours, and selective colour result."
)
add_labeled_paragraph(
    doc,
    "What the figure shows",
    f"White mask pixels passed all four tests and black pixels failed at least one. The mask selects "
    f"{METRICS['task1']['selected_percent']:.2f}% of the image. I use bitwise_and to take original colour under "
    "the mask and grayscale under its inverse, then bitwise_or to join the nonoverlapping parts. Red skin stays "
    "coloured, while yellow-green mottling, highlights, stalks, and shadowed regions inside the same apples can "
    "turn gray. The remaining coloured ribbon strip and the gray holes show the limitation: this rule identifies "
    "red-dominant pixels, not complete apples. Whole-object detection would need spatial or shape information."
)

add_task_heading(doc, "Task 2 Creating a Ghostly Overlay")
add_labeled_paragraph(
    doc,
    "Preparing the images",
    "Weighted addition works at corresponding pixel positions, so both arrays need the same height, width, channels, "
    "and data type. The cameraman is 225 x 225. I resized the zebra proportionally from 425 x 313 to 225 x 166, "
    "then added 29 white rows above and 30 below. This keeps the whole zebra without stretching it. Both prepared "
    "inputs are 225 x 225 uint8 BGR images on a common canvas, although their unrelated scenes are not registered."
)
add_native_equation(doc, "I = alpha C + beta Z + gamma")
doc.add_paragraph(
    "I used gamma = 0 and weights that sum to 1. Alpha is the cameraman contribution and beta is the zebra "
    "contribution at the same pixel. For example, equal weights combine intensities 80 and 240 as about 160 after rounding."
)
t2 = doc.add_table(rows=4, cols=3)
for j, text in enumerate(("Experiment", "Alpha cameraman", "Beta zebra")):
    t2.cell(0, j).text = text
rows = (("Cameraman dominant", "0.75", "0.25"), ("Equal contribution", "0.50", "0.50"),
        ("Zebra dominant", "0.25", "0.75"))
for i, row in enumerate(rows, start=1):
    for j, text in enumerate(row):
        t2.cell(i, j).text = text
style_table(t2, (Inches(3.4), Inches(1.75), Inches(1.65)), numeric_columns=(1, 2))
add_figure(
    doc,
    "task2/comparison.png",
    "Figure 2  Prepared inputs and three pixel by pixel weighted blends",
    width=Inches(5.75),
    alt_text="Cameraman and zebra inputs followed by blends with weights 0.75 and 0.25, 0.50 and 0.50, and 0.25 and 0.75."
)
add_labeled_paragraph(
    doc,
    "What the figure shows",
    "At 0.75 and 0.25 the cameraman and tripod dominate while the zebra is faint. Equal weights give the two "
    "silhouettes similar strength. At 0.25 and 0.75 the zebra dominates and the cameraman looks pale behind it. "
    "The frame also becomes brighter as beta increases because most zebra background pixels are white and are "
    "mixed into the whole cameraman scene. The ghostly look comes from superimposing two semi-weighted subjects. "
    "A cleaner composite would require foreground extraction, which is outside this full-image blending task."
)

add_task_heading(doc, "Task 3 Artistic Vignette Filtering")
add_labeled_paragraph(
    doc,
    "Reasoning and method",
    "I wanted the rover and nearby terrain to remain visible while the frame edges became darker. I drew a filled "
    f"white circle of radius {METRICS['task3']['radius']} pixels on a black one-channel mask. White represents a "
    "multiplication weight near 1 and black a weight near 0. Applying the hard circle directly would create an "
    "obvious boundary, so I blurred it with a Gaussian filter. Sigma is "
    f"{METRICS['task3']['sigma']:.1f} pixels and the kernel is {METRICS['task3']['kernel_size']} x "
    f"{METRICS['task3']['kernel_size']}. After normalization, the floating-point mask contains continuous weights from 0 to 1."
)
add_native_equation(doc, "V(y,x,c) = round(I(y,x,c) x W(y,x))")
doc.add_paragraph(
    "A weight of 1 preserves a source value, 0.5 halves it, and a value near 0 makes it almost black. I apply the "
    "same weight to B, G, and R, so the intended change is brightness rather than hue."
)
add_figure(
    doc,
    "task3/stages.png",
    "Figure 3  Hard circle smooth normalized weights and the resulting vignette",
    width=Inches(6.25),
    alt_text="Original Mars rover image, hard circular mask, smooth normalized mask, and final vignette."
)
t3 = METRICS["task3"]
table3 = doc.add_table(rows=3, cols=3)
for j, text in enumerate(("Measured region", "Original mean", "Vignette mean")):
    table3.cell(0, j).text = text
values = (
    ("Centre quarter grayscale", f"{t3['center_before_mean']:.2f}", f"{t3['center_after_mean']:.2f}"),
    ("Four corner patches grayscale", f"{t3['corner_before_mean']:.2f}", f"{t3['corner_after_mean']:.2f}"),
)
for i, row in enumerate(values, start=1):
    for j, text in enumerate(row):
        table3.cell(i, j).text = text
style_table(table3, (Inches(4.0), Inches(1.4), Inches(1.4)), numeric_columns=(1, 2))
add_labeled_paragraph(
    doc,
    "What the figure shows",
    f"The centre pixel is unchanged and the centre-quarter mean changes from {t3['center_before_mean']:.2f} to "
    f"{t3['center_after_mean']:.2f}; the corner mean falls from {t3['corner_before_mean']:.2f} to "
    f"{t3['corner_after_mean']:.2f}. The rover body, wheels, arm, and central rock texture remain readable. A "
    "smaller radius would darken more of the rover, while a larger one would retain more of the frame. More blur "
    "would spread the transition and less blur would reveal the circle. A nonzero minimum weight could preserve "
    "corner detail, but the current near-black corners give the requested strong framing."
)

add_task_heading(doc, "Task 4 Channel Swapping")
add_labeled_paragraph(
    doc,
    "Reasoning and method",
    "Each colour pixel contains blue, green, and red intensities. OpenCV stores them in BGR order. Splitting the "
    "image exposes those planes; merging them in a new order changes which display colour receives each original "
    "plane. Output BGR (B,R,G), for example, keeps blue, displays original red as green, and displays original green "
    "as red. A neutral pixel has B = G = R, so any permutation of its three equal values is unchanged."
)
add_figure(
    doc,
    "task4/comparison.png",
    "Figure 4  Original zebra and three exact channel permutations",
    width=Inches(6.25),
    alt_text="Original zebra and output BGR arrangements R G B, B R G, and G B R."
)
zebra_rows = METRICS["task4"]["experiments"][:3]
table4 = doc.add_table(rows=4, cols=4)
for j, text in enumerate(("Output BGR", "Mean change", "Maximum", "Changed pixels")):
    table4.cell(0, j).text = text
for i, row in enumerate(zebra_rows, start=1):
    values = (row["BGR_tuple"], f"{row['MAE']:.3f}", str(row["max_difference"]),
              f"{row['changed_pixel_percent']:.2f}%")
    for j, text in enumerate(values):
        table4.cell(i, j).text = text
style_table(table4, (Inches(2.0), Inches(1.5), Inches(1.35), Inches(1.6)), numeric_columns=(1, 2, 3))
add_labeled_paragraph(
    doc,
    "What the figure shows",
    f"The ears, muzzle, tail, and pale shadow shift colour, while the white background and black stripes remain "
    f"almost unchanged. Exactly {METRICS['task4']['neutral_zebra_percent']:.2f}% of pixels have B = G = R. "
    "About 45 to 51% of pixels change, but the mean differences are only 1.41 to 2.36 levels out of 255 because "
    "many channel values differ only slightly. The experiment changes colour interpretation, not image geometry, "
    "and explains why a large changed-pixel count can still look subtle."
)

add_task_heading(doc, "Channel Swapping on the Mars Image")
doc.add_paragraph(
    "I repeated the same assignments on the Mars image without adding a tint or changing contrast. The orange-brown "
    "terrain has more unequal BGR values, so this supplementary example makes the channel effects easier to see."
)
add_figure(
    doc,
    "task4/supplementary_marsrover.png",
    "Figure 5  The same channel permutations on the Mars image",
    width=Inches(6.25),
    alt_text="Original Mars rover image and output BGR arrangements R G B, B R G, and G B R."
)
mars_rows = METRICS["task4"]["experiments"][3:]
table5 = doc.add_table(rows=4, cols=3)
for j, text in enumerate(("Output BGR", "Zebra mean change", "Mars mean change")):
    table5.cell(0, j).text = text
for i, (zebra_row, mars_row) in enumerate(zip(zebra_rows, mars_rows), start=1):
    values = (zebra_row["BGR_tuple"], f"{zebra_row['MAE']:.3f}", f"{mars_row['MAE']:.3f}")
    for j, text in enumerate(values):
        table5.cell(i, j).text = text
style_table(table5, (Inches(2.7), Inches(2.0), Inches(2.0)), numeric_columns=(1, 2))
add_labeled_paragraph(
    doc,
    "What the figure shows",
    "The three Mars outputs are strongly blue, green, and magenta. Their mean changes are 29.796, 18.264, and "
    "11.737 levels, much larger than for the zebra. Orange terrain originally has more red than blue, so sending "
    "those unequal values to different display channels changes most of the frame. Neutral rover panels remain "
    "more stable. The comparison shows that the strength of a permutation depends on the source channel differences."
)
conclusion_heading = doc.add_heading("What I learned and what remains limited", level=2)
conclusion_heading.paragraph_format.page_break_before = True
doc.add_paragraph(
    "Complementary masks combine colour and grayscale exactly, but a colour threshold does not understand object "
    "boundaries. Weighted addition makes both images visible while also mixing their backgrounds. Gaussian smoothing "
    "turns a binary circle into controllable brightness weights, although the framing depends on radius and blur. "
    "Channel permutation preserves spatial detail while changing colour meaning, and neutral pixels explain why some "
    f"sources respond less than others. The executed notebook uses these same parameters and reopens all {DATA['png_count']} saved PNG files."
)
source = doc.add_paragraph()
lead = source.add_run("Source ")
lead.bold = True
source.add_run("CSE480 CSE482 Lab Assignment 01 supplied with the project")
repository = doc.add_paragraph()
lead = repository.add_run("Project repository ")
lead.bold = True
add_hyperlink(repository, REPOSITORY_URL, REPOSITORY_URL)

doc.core_properties.title = "Computer Vision Lab Assignment 01"
doc.core_properties.subject = "CSE480 CSE482 Lab Report"
doc.core_properties.author = "Nada Said Yassin Morgan"
doc.core_properties.keywords = "OpenCV, computer vision, selective colour, blending, vignette, channel swapping"
doc.core_properties.comments = "Verified lab report generated from the executed notebook results."

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
